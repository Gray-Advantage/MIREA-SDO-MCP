"""Постраничное чтение файлов СДО без скачивания их пользователю.

Файл кладётся во внутренний кэш, из него извлекается текст только запрошенных
страниц. Что считать «страницей», зависит от формата: у PDF это настоящие
страницы, у презентации — слайды, у таблицы — листы, у текстовых форматов —
куски примерно одинакового размера.
"""

from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Any, Callable

from . import config
from .client import SdoClient, SdoError, absolute

#: Размер синтетической «страницы» для форматов без собственного разбиения.
CHUNK_CHARS = 3000

TEXT_SUFFIXES = {".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".html", ".htm", ".py", ".log"}


class UnsupportedFormat(SdoError):
    """Формат файла не поддерживается для чтения текстом."""


async def read_file(
    client: SdoClient,
    url: str | None = None,
    path: str | None = None,
    pages: str | None = None,
    max_chars: int = 20000,
) -> dict[str, Any]:
    """Читает файл постранично. Источник — ссылка в СДО либо локальный путь."""
    if not url and not path:
        raise SdoError("Укажи url (ссылку в СДО) или path (локальный файл).")

    if path:
        local = Path(path).expanduser()
        if not local.exists():
            raise SdoError(f"Файл не найден: {local}")
        source = str(local)
    else:
        local = await _fetch_to_cache(client, absolute(url))
        source = url

    extractor = _pick_extractor(local)
    all_pages = extractor(local)
    total = len(all_pages)
    wanted = _parse_pages(pages, total)

    chunks: list[str] = []
    used: list[int] = []
    empty: list[int] = []
    truncated = False
    length = 0
    for number in wanted:
        body = all_pages[number - 1]
        if not body:
            # Страница без извлекаемого текста — скорее всего скан или картинка.
            empty.append(number)
            continue
        if length + len(body) > max_chars:
            body = body[: max(0, max_chars - length)]
            truncated = True
        chunks.append(f"--- стр. {number} ---\n{body}")
        used.append(number)
        length += len(body)
        if truncated:
            break

    return {
        "source": source,
        "name": local.name,
        "format": local.suffix.lower().lstrip("."),
        "total_pages": total,
        "pages_requested": wanted,
        "pages_returned": used,
        "empty_pages": empty,
        "truncated": truncated,
        "size_bytes": local.stat().st_size,
        "text": "\n\n".join(chunks),
        "note": (
            f"Страницы {empty} без текстового слоя (вероятно, изображения)."
            if empty
            else None
        ),
    }


async def _fetch_to_cache(client: SdoClient, url: str) -> Path:
    digest = hashlib.sha256(url.encode("utf-8")).hexdigest()[:16]
    directory = config.CACHE_DIR / digest
    if directory.is_dir():
        existing = [p for p in directory.iterdir() if p.is_file() and p.suffix != ".part"]
        if existing:
            return existing[0]
    result = await client.download(url, directory)
    return Path(result["path"])


# --- разбиение по страницам -------------------------------------------


def _parse_pages(spec: str | None, total: int) -> list[int]:
    """Разбирает ``"3"``, ``"1-5"``, ``"1,4,7-9"``; без указания — первые пять."""
    if total == 0:
        return []
    if not spec:
        return list(range(1, min(total, 5) + 1))

    wanted: list[int] = []
    for part in re.split(r"[,\s]+", spec.strip()):
        if not part:
            continue
        match = re.fullmatch(r"(\d+)\s*-\s*(\d+)", part)
        if match:
            start, end = int(match.group(1)), int(match.group(2))
            wanted.extend(range(min(start, end), max(start, end) + 1))
        elif part.isdigit():
            wanted.append(int(part))
        else:
            raise SdoError(f"Не понимаю диапазон страниц: {part!r}. Примеры: 3, 1-5, 1,4,7-9")

    seen: set[int] = set()
    return [p for p in wanted if 1 <= p <= total and not (p in seen or seen.add(p))]


def _pick_extractor(path: Path) -> Callable[[Path], list[str]]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _pdf_pages
    if suffix == ".docx":
        return _docx_pages
    if suffix == ".pptx":
        return _pptx_pages
    if suffix in {".xlsx", ".xlsm"}:
        return _xlsx_pages
    if suffix in TEXT_SUFFIXES:
        return _text_pages
    raise UnsupportedFormat(
        f"Чтение файлов «{suffix or 'без расширения'}» не поддерживается. "
        "Скачай его через download_file и открой сам. "
        f"Поддерживаются: pdf, docx, pptx, xlsx, {', '.join(sorted(s.lstrip('.') for s in TEXT_SUFFIXES))}."
    )


def _pdf_pages(path: Path) -> list[str]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return [(page.extract_text() or "").strip() for page in reader.pages]


def _docx_pages(path: Path) -> list[str]:
    import docx

    document = docx.Document(str(path))
    blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return _chunk(blocks)


def _pptx_pages(path: Path) -> list[str]:
    from pptx import Presentation

    slides: list[str] = []
    for index, slide in enumerate(Presentation(str(path)).slides, start=1):
        lines = [f"[слайд {index}]"]
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
        slides.append("\n".join(lines))
    return slides


def _xlsx_pages(path: Path) -> list[str]:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    sheets: list[str] = []
    for sheet in workbook.worksheets:
        lines = [f"[лист {sheet.title}]"]
        for row in sheet.iter_rows(values_only=True):
            values = ["" if v is None else str(v) for v in row]
            if any(values):
                lines.append(" | ".join(values))
        sheets.append("\n".join(lines))
    workbook.close()
    return sheets


def _text_pages(path: Path) -> list[str]:
    content = path.read_text("utf-8", errors="replace")
    return _chunk(content.splitlines())


def _chunk(lines: list[str]) -> list[str]:
    """Склеивает строки в куски примерно по CHUNK_CHARS символов."""
    pages: list[str] = []
    current: list[str] = []
    size = 0
    for line in lines:
        if size + len(line) > CHUNK_CHARS and current:
            pages.append("\n".join(current))
            current, size = [], 0
        current.append(line)
        size += len(line) + 1
    if current:
        pages.append("\n".join(current))
    return pages or [""]
